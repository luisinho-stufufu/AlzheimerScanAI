"""Create subject-level binary CDR splits and a linked local image view."""

from __future__ import annotations

import argparse
import csv
import json
import os
import random
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_METADATA = PROJECT_ROOT / "data/raw/oasis-cross-sectional-5708aa0a98d82080.xlsx"
DEFAULT_IMAGE_ROOT = PROJECT_ROOT / "data/raw/images"
DEFAULT_OUTPUT = PROJECT_ROOT / "data/processed/binary_cdr"
SPLIT_FRACTIONS = (("train", 0.70), ("validation", 0.15), ("test", 0.15))
SPLIT_COLUMNS = (
    "participant_id",
    "session_id",
    "cdr",
    "binary_label",
    "class_name",
    "source_image_path",
    "organized_image_path",
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--metadata", type=Path, default=DEFAULT_METADATA)
    parser.add_argument("--image-root", type=Path, default=DEFAULT_IMAGE_ROOT)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--seed", type=int, default=42)
    return parser.parse_args()


def cdr_value(value: Any) -> float | None:
    if value is None or (isinstance(value, str) and value.strip().upper() in {"", "N/A", "NA"}):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Unexpected CDR value: {value!r}") from exc
    if number not in {0.0, 0.5, 1.0, 2.0, 3.0}:
        raise ValueError(f"Unexpected CDR value: {value!r}")
    return number


def load_baseline_rows(metadata: Path) -> list[dict[str, Any]]:
    if not metadata.is_file():
        raise FileNotFoundError(f"Metadata workbook not found: {metadata}")
    workbook = load_workbook(metadata, read_only=True, data_only=True)
    try:
        if "oasis_cross-sectional" not in workbook.sheetnames:
            raise ValueError("Workbook is missing the 'oasis_cross-sectional' sheet.")
        sheet = workbook["oasis_cross-sectional"]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if headers is None:
            raise ValueError("Metadata worksheet is empty.")
        columns = {str(name).strip(): index for index, name in enumerate(headers) if name is not None}
        if not {"ID", "CDR"}.issubset(columns):
            raise ValueError("Metadata worksheet must contain ID and CDR columns.")

        records: list[dict[str, Any]] = []
        seen_participants: set[str] = set()
        for row in rows:
            raw_session = row[columns["ID"]]
            if raw_session is None:
                continue
            session_id = str(raw_session).strip()
            if not session_id.endswith("_MR1"):
                continue
            participant_id = session_id.removesuffix("_MR1")
            if participant_id in seen_participants:
                raise ValueError(f"Duplicate baseline record for {participant_id}.")
            seen_participants.add(participant_id)
            records.append(
                {
                    "participant_id": participant_id,
                    "session_id": session_id,
                    "cdr": cdr_value(row[columns["CDR"]]),
                }
            )
        return records
    finally:
        workbook.close()


def index_t1_volumes(image_root: Path) -> dict[str, Path]:
    if not image_root.is_dir():
        raise FileNotFoundError(f"MRI image directory not found: {image_root}")
    volumes: dict[str, Path] = {}
    pattern = "oasis_cs_freesurfer_disc*/disc*/OAS1_*_MR1/mri/T1.mgz"
    for volume in image_root.glob(pattern):
        participant_id = volume.parents[1].name.removesuffix("_MR1")
        if participant_id in volumes:
            raise ValueError(f"Multiple processed T1 volumes found for {participant_id}.")
        volumes[participant_id] = volume
    if not volumes:
        raise FileNotFoundError(f"No processed T1.mgz volumes found under {image_root}.")
    return volumes


def stratified_assignments(
    records: list[dict[str, Any]], seed: int
) -> dict[str, str]:
    by_label: dict[int, list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        by_label[record["binary_label"]].append(record)

    assignments: dict[str, str] = {}
    for label, members in sorted(by_label.items()):
        random.Random(seed + label).shuffle(members)
        exact_counts = [len(members) * fraction for _, fraction in SPLIT_FRACTIONS]
        counts = [int(count) for count in exact_counts]
        remainder = len(members) - sum(counts)
        ranked_splits = sorted(
            range(len(counts)),
            key=lambda index: (exact_counts[index] - counts[index], -index),
            reverse=True,
        )
        for index in ranked_splits[:remainder]:
            counts[index] += 1

        start = 0
        for (split, _), count in zip(SPLIT_FRACTIONS, counts):
            for record in members[start : start + count]:
                assignments[record["participant_id"]] = split
            start += count
    return assignments


def create_link(target: Path, link: Path) -> None:
    link.parent.mkdir(parents=True, exist_ok=True)
    relative_target = os.path.relpath(target, link.parent)
    if link.is_symlink():
        if link.resolve() != target.resolve():
            raise FileExistsError(f"Existing image link points elsewhere: {link}")
        return
    if link.exists():
        raise FileExistsError(f"Refusing to replace existing file: {link}")
    link.symlink_to(relative_target)


def write_csv(path: Path, fieldnames: tuple[str, ...], rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    args = parse_args()
    baseline = load_baseline_rows(args.metadata)
    volumes = index_t1_volumes(args.image_root)

    included: list[dict[str, Any]] = []
    excluded: list[dict[str, Any]] = []
    for record in baseline:
        cdr = record["cdr"]
        if cdr is None:
            excluded.append({**record, "exclusion_reason": "missing_cdr"})
            continue
        if record["participant_id"] not in volumes:
            excluded.append({**record, "exclusion_reason": "missing_local_t1_volume"})
            continue
        label = int(cdr > 0)
        included.append({**record, "binary_label": label, "class_name": "cdr_gt_0" if label else "cdr_0"})

    if not included:
        raise ValueError("No labeled subjects with an available processed T1 volume were found.")

    assignments = stratified_assignments(included, args.seed)
    if len(assignments) != len(included):
        raise RuntimeError("Split assignment did not include every eligible subject.")

    output_dir = args.output_dir
    manifest_rows: dict[str, list[dict[str, Any]]] = {split: [] for split, _ in SPLIT_FRACTIONS}
    split_counts: dict[str, Counter[str]] = {split: Counter() for split, _ in SPLIT_FRACTIONS}
    image_root = args.image_root
    for record in included:
        participant_id = record["participant_id"]
        split = assignments[participant_id]
        volume = volumes[participant_id]
        relative_source = volume.relative_to(image_root)
        organized = (
            output_dir
            / "images"
            / split
            / record["class_name"]
            / participant_id
            / "T1.mgz"
        )
        create_link(volume, organized)
        source_path = Path(os.path.relpath(image_root / relative_source, PROJECT_ROOT))
        organized_path = Path(os.path.relpath(organized, PROJECT_ROOT))
        manifest_rows[split].append(
            {
                "participant_id": participant_id,
                "session_id": record["session_id"],
                "cdr": record["cdr"],
                "binary_label": record["binary_label"],
                "class_name": record["class_name"],
                "source_image_path": source_path.as_posix(),
                "organized_image_path": organized_path.as_posix(),
            }
        )
        split_counts[split][record["class_name"]] += 1

    for split, rows in manifest_rows.items():
        rows.sort(key=lambda row: row["participant_id"])
        write_csv(output_dir / "splits" / f"{split}.csv", SPLIT_COLUMNS, rows)

    write_csv(
        output_dir / "splits" / "excluded.csv",
        ("participant_id", "session_id", "cdr", "exclusion_reason"),
        sorted(excluded, key=lambda row: row["participant_id"]),
    )
    summary = {
        "task": "binary CDR classification (CDR 0 vs CDR > 0)",
        "seed": args.seed,
        "split_fractions": dict(SPLIT_FRACTIONS),
        "baseline_subjects": len(baseline),
        "eligible_subjects": len(included),
        "excluded_missing_cdr": sum(row["exclusion_reason"] == "missing_cdr" for row in excluded),
        "excluded_missing_local_t1_volume": sum(
            row["exclusion_reason"] == "missing_local_t1_volume" for row in excluded
        ),
        "class_totals": dict(Counter(record["class_name"] for record in included)),
        "split_class_counts": {
            split: dict(counts) for split, counts in split_counts.items()
        },
    }
    summary_path = output_dir / "splits" / "summary.json"
    summary_path.parent.mkdir(parents=True, exist_ok=True)
    summary_path.write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")

    print(f"Eligible unique subjects: {len(included)}")
    print(f"Excluded for missing CDR: {summary['excluded_missing_cdr']}")
    print(f"Excluded for missing local T1 volume: {summary['excluded_missing_local_t1_volume']}")
    for split, counts in split_counts.items():
        print(f"{split}: CDR 0={counts['cdr_0']}, CDR > 0={counts['cdr_gt_0']}, total={sum(counts.values())}")
    print(f"Local subject-level image view and split manifests: {output_dir}")


if __name__ == "__main__":
    main()
